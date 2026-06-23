from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "ordinary-user-manual.md"
OUTPUT = ROOT / "docs" / "ReviewIQ-普通用户操作手册.docx"
PUBLIC_OUTPUT = ROOT / "apps" / "web" / "public" / "downloads" / "ReviewIQ-ordinary-user-manual.docx"


def set_east_asia_font(style, font_name: str) -> None:
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_run_font(run, font_name: str = "Microsoft YaHei") -> None:
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr)
    run._r.append(fld_char_2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    set_east_asia_font(normal, "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = doc.styles[name]
        set_east_asia_font(style, "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    for name in ["List Bullet", "List Number"]:
        style = doc.styles[name]
        set_east_asia_font(style, "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "ReviewIQ 普通用户操作手册"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if header.runs:
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor(100, 116, 139)
        set_run_font(header.runs[0])

    footer = section.footer.paragraphs[0]
    run = footer.add_run("第 ")
    set_run_font(run)
    add_page_number(footer)


def add_cover(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(70)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("ReviewIQ 普通用户操作手册")
    set_run_font(run)
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("评论采集、持续监听、AI 分析、报告阅读与评论复核操作指南")
    set_run_font(run)
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(71, 85, 105)

    meta = doc.add_paragraph()
    run = meta.add_run("版本：2026-06-23\n适用对象：普通业务用户、运营人员、内容分析人员")
    set_run_font(run)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(71, 85, 105)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    note.paragraph_format.space_after = Pt(12)
    run = note.add_run(
        "普通用户不需要配置模型、提示词、抓取参数或系统后台。平台会统一使用管理员或超管维护好的 AI 分析配置和评论采集配置。"
    )
    set_run_font(run)
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_page_break()


def add_toc(doc: Document, headings: list[tuple[int, str]]) -> None:
    p = doc.add_paragraph("目录", style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    for level, text in headings:
        if level > 2:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25 if level == 2 else 0)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        set_run_font(run)
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(51, 65, 85)
    doc.add_page_break()


def add_image(doc: Document, alt: str, rel_path: str) -> None:
    image_path = SOURCE.parent / rel_path
    if not image_path.exists():
        p = doc.add_paragraph(f"[图片缺失：{rel_path}]")
        p.runs[0].font.color.rgb = RGBColor(155, 28, 28)
        return

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    run = caption.add_run(f"图：{alt}")
    set_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(71, 85, 105)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    width = Inches(2.6 if "mobile" in image_path.name else 6.35)
    run.add_picture(str(image_path), width=width)


def clean_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    return text.strip()


def add_body_text(doc: Document, line: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(clean_inline(line))
    set_run_font(run)


def build() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    headings: list[tuple[int, str]] = []
    for line in lines:
        match = re.match(r"^(#{1,3})\s+(.+)$", line)
        if match:
            level = len(match.group(1))
            title = clean_inline(match.group(2))
            if level <= 2:
                headings.append((level, title))

    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_toc(doc, headings)

    for line in lines:
        if not line.strip():
            continue

        image = re.match(r"^!\[(.+?)\]\((.+?)\)$", line.strip())
        if image:
            add_image(doc, image.group(1), image.group(2))
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = min(len(heading.group(1)), 3)
            text = clean_inline(heading.group(2))
            doc.add_paragraph(text, style=f"Heading {level}")
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            p = doc.add_paragraph(clean_inline(numbered.group(1)), style="List Number")
            for run in p.runs:
                set_run_font(run)
            continue

        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(clean_inline(bullet.group(1)), style="List Bullet")
            for run in p.runs:
                set_run_font(run)
            continue

        add_body_text(doc, line)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.name is None:
                set_run_font(run)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    PUBLIC_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    shutil.copy2(OUTPUT, PUBLIC_OUTPUT)
    print(OUTPUT)
    print(PUBLIC_OUTPUT)


if __name__ == "__main__":
    build()
